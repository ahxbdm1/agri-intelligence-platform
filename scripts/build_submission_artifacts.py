"""从竞赛 Markdown 源生成 Word 与 PDF 提交文档。

文件名统一使用赛题通知规定的 团队名称_学校名称_队长姓名_材料名 格式，
当前团队信息为：云穗智擎 / 山东建筑大学 / 队长 张衡。
PPT 由 scripts/build_pptx.py 单独生成。
"""

from __future__ import annotations

from html import escape
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = ROOT / "submission_artifacts"
BLACK = RGBColor(31, 31, 31)
DARK_GRAY = RGBColor(68, 68, 68)
MID_GRAY = RGBColor(102, 102, 102)
LIGHT_GRAY_HEX = "F2F2F2"
BORDER_GRAY_HEX = "B7B7B7"


def clean_inline(text: str) -> str:
    """Remove Markdown-only inline markers from formal office documents."""
    text = re.sub(r"`([^`]*)`", r"\1", text)
    return re.sub(r"\*\*([^*]+)\*\*", r"\1", text)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    run.font.name = "SimSun"
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)
    paragraph.add_run(" 页")


def configure_doc(document: Document, title: str) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    normal = document.styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color in (("Heading 1", 16, BLACK), ("Heading 2", 12.5, DARK_GRAY), ("Heading 3", 11, DARK_GRAY)):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "农智云瞰项目申报材料"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Microsoft YaHei"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = MID_GRAY
    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.name = "SimSun"
        run.font.size = Pt(8)
        run.font.color.rgb = MID_GRAY

    title_para = document.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(title)
    title_run.font.name = "Microsoft YaHei"
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = BLACK
    title_para.paragraph_format.space_after = Pt(4)
    subtitle = document.add_paragraph("山东省大学生软件设计大赛 · 大数据与人工智能行业应用开发赛道 参赛材料　|　云穗智擎 · 山东建筑大学 · 队长 张衡")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(9)
    subtitle.runs[0].font.color.rgb = MID_GRAY
    document.add_paragraph()


def add_markdown_to_docx(document: Document, source: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_rows
        if not table_rows:
            return
        rows = [row for row in table_rows if not all(re.fullmatch(r"\s*:?-+:?\s*", cell) for cell in row)]
        if rows:
            table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
            table.style = "Table Grid"
            for r_index, row in enumerate(rows):
                for c_index in range(len(table.columns)):
                    cell = table.cell(r_index, c_index)
                    set_cell_margins(cell)
                    cell.text = clean_inline(row[c_index]) if c_index < len(row) else ""
                    if r_index == 0:
                        set_cell_shading(cell, LIGHT_GRAY_HEX)
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
                            run.font.color.rgb = BLACK
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.name = "SimSun"
                            run.font.size = Pt(9)
        table_rows = []

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            table_rows.append([part.strip() for part in stripped.strip("|").split("|")])
            continue
        flush_table()
        if not stripped:
            continue
        if stripped.startswith("# "):
            continue
        if stripped.startswith("## "):
            document.add_heading(stripped[3:].strip(), level=1)
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:].strip(), level=2)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(clean_inline(stripped[2:].strip()))
        elif re.match(r"^\d+\.\s", stripped):
            paragraph = document.add_paragraph(style="List Number")
            paragraph.add_run(clean_inline(re.sub(r"^\d+\.\s", "", stripped)))
        elif stripped.startswith("```"):
            continue
        else:
            document.add_paragraph(clean_inline(stripped))
    flush_table()


def build_docx(source_name: str, output_name: str, title: str) -> Path:
    document = Document()
    configure_doc(document, title)
    add_markdown_to_docx(document, DOCS / source_name)
    target = OUT / output_name
    document.save(target)
    return target


def register_pdf_font() -> str:
    """按平台依次寻找可用的中文字体，Windows 与 Linux 均可生成 PDF。"""
    candidates = [
        Path("C:/Windows/Fonts/NotoSansSC-VF.ttf"),
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"),
    ]
    for font_path in candidates:
        if not font_path.exists():
            continue
        try:
            if font_path.suffix.lower() == ".ttc":
                pdfmetrics.registerFont(TTFont("AgriSans", str(font_path), subfontIndex=0))
            else:
                pdfmetrics.registerFont(TTFont("AgriSans", str(font_path)))
            return "AgriSans"
        except Exception:  # noqa: BLE001  字体不可用则继续尝试下一个
            continue
    raise FileNotFoundError("未找到可用的中文字体，请安装 Noto Sans CJK 或微软雅黑后重试")


def pdf_styles(font_name: str):
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("AgriTitle", parent=styles["Title"], fontName=font_name, fontSize=19, leading=25, textColor=colors.HexColor("#1F1F1F"), alignment=TA_CENTER, spaceAfter=4),
        "subtitle": ParagraphStyle("AgriSubtitle", parent=styles["Normal"], fontName=font_name, fontSize=8.5, leading=12, textColor=colors.HexColor("#666666"), alignment=TA_CENTER, spaceAfter=15),
        "h1": ParagraphStyle("AgriH1", parent=styles["Heading1"], fontName=font_name, fontSize=14, leading=19, textColor=colors.HexColor("#1F1F1F"), spaceBefore=10, spaceAfter=6),
        "h2": ParagraphStyle("AgriH2", parent=styles["Heading2"], fontName=font_name, fontSize=11.5, leading=16, textColor=colors.HexColor("#444444"), spaceBefore=8, spaceAfter=4),
        "body": ParagraphStyle("AgriBody", parent=styles["BodyText"], fontName=font_name, fontSize=9.5, leading=15, textColor=colors.HexColor("#1F1F1F"), alignment=TA_LEFT, spaceAfter=5),
        "bullet": ParagraphStyle("AgriBullet", parent=styles["BodyText"], fontName=font_name, fontSize=9.5, leading=15, leftIndent=12, firstLineIndent=-8, textColor=colors.HexColor("#1F1F1F"), spaceAfter=3),
    }


def build_pdf(source_name: str, output_name: str, title: str) -> Path:
    font_name = register_pdf_font()
    styles = pdf_styles(font_name)
    story = [Paragraph(escape(title), styles["title"]), Paragraph("山东省大学生软件设计大赛 · 大数据与人工智能行业应用开发赛道 参赛材料　|　云穗智擎 · 山东建筑大学 · 队长 张衡", styles["subtitle"])]
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_rows
        if not table_rows:
            return
        rows = [row for row in table_rows if not all(re.fullmatch(r"\s*:?-+:?\s*", cell) for cell in row)]
        if rows:
            normalized = [[Paragraph(escape(clean_inline(cell)), styles["body"]) for cell in row] for row in rows]
            table = Table(normalized, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F2F2")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1F1F1F")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B7B7B7")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.extend([Spacer(1, 3), table, Spacer(1, 7)])
        table_rows = []

    for line in (DOCS / source_name).read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            table_rows.append([part.strip() for part in stripped.strip("|").split("|")])
            continue
        flush_table()
        if not stripped or stripped.startswith("# ") or stripped.startswith("```"):
            continue
        if stripped.startswith("## "):
            story.append(Paragraph(escape(stripped[3:].strip()), styles["h1"]))
        elif stripped.startswith("### "):
            story.append(Paragraph(escape(stripped[4:].strip()), styles["h2"]))
        elif stripped.startswith("- ") or stripped.startswith("* "):
            story.append(Paragraph("• " + escape(clean_inline(stripped[2:].strip())), styles["bullet"]))
        elif re.match(r"^\d+\.\s", stripped):
            story.append(Paragraph(escape(clean_inline(stripped)), styles["bullet"]))
        else:
            story.append(Paragraph(escape(clean_inline(stripped)), styles["body"]))
    flush_table()

    target = OUT / output_name
    document = SimpleDocTemplate(str(target), pagesize=A4, rightMargin=20 * mm, leftMargin=20 * mm, topMargin=18 * mm, bottomMargin=18 * mm, title=title, author="农智云瞰项目组")
    document.build(story)
    return target


def main() -> None:
    OUT.mkdir(exist_ok=True)
    outputs = [
        build_docx("项目概要介绍.md", "云穗智擎_山东建筑大学_张衡_项目概要介绍.docx", "项目概要介绍"),
        build_docx("项目详细方案.md", "云穗智擎_山东建筑大学_张衡_项目详细方案.docx", "项目详细方案"),
        build_docx("需求分析说明书.md", "云穗智擎_山东建筑大学_张衡_需求分析说明书.docx", "需求分析说明书"),
        build_docx("数据库设计说明书.md", "云穗智擎_山东建筑大学_张衡_数据库设计说明书.docx", "数据库设计说明书"),
        build_docx("概要设计说明书.md", "云穗智擎_山东建筑大学_张衡_概要设计说明书.docx", "概要设计说明书"),
        build_docx("详细设计说明书.md", "云穗智擎_山东建筑大学_张衡_详细设计说明书.docx", "详细设计说明书"),
        build_docx("安装部署说明.md", "云穗智擎_山东建筑大学_张衡_安装部署说明书.docx", "安装部署说明书"),
        build_docx("使用说明.md", "云穗智擎_山东建筑大学_张衡_软件使用说明书.docx", "软件使用说明书"),
        build_docx("测试报告.md", "云穗智擎_山东建筑大学_张衡_测试文档.docx", "测试文档"),
        build_docx("工作计划.md", "云穗智擎_山东建筑大学_张衡_工作计划.docx", "工作计划"),
        build_docx("会议纪要.md", "云穗智擎_山东建筑大学_张衡_会议纪要.docx", "会议纪要"),
        build_docx("工作总结.md", "云穗智擎_山东建筑大学_张衡_工作总结.docx", "工作总结"),
        build_docx("作品基本信息登记表.md", "云穗智擎_山东建筑大学_张衡_作品基本信息登记表.docx", "作品基本信息登记表"),
        build_docx("数据集说明.md", "云穗智擎_山东建筑大学_张衡_数据集说明.docx", "数据集说明"),
        build_docx("开源组件与协议说明.md", "云穗智擎_山东建筑大学_张衡_开源组件与协议说明.docx", "开源组件与协议说明"),
        build_docx("竞赛承诺书.md", "云穗智擎_山东建筑大学_张衡_竞赛承诺书.docx", "竞赛承诺书"),
        build_pdf("项目概要介绍.md", "云穗智擎_山东建筑大学_张衡_项目概要介绍.pdf", "项目概要介绍"),
        build_pdf("项目详细方案.md", "云穗智擎_山东建筑大学_张衡_项目详细方案.pdf", "项目详细方案"),
    ]
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
